"""Report parser for text-like reports: .txt, .md and .docx. (PDF reports go through scan.py.)"""
import io
import re

from .common import IngestError, ParseResult, make_unit

MIN_PARAGRAPH_CHARS = 2
RE_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")


def _section(path):
    return " > ".join(h for h in path if h)


def parse_text_report(text: str) -> ParseResult:
    result = ParseResult()
    headings = {}
    block, start = [], 0

    def flush(end_line):
        nonlocal block
        body = " ".join(line.strip() for line in block).strip()
        if len(body) >= MIN_PARAGRAPH_CHARS:
            path = [headings[k] for k in sorted(headings)]
            result.units.append(make_unit(body, {"line": start, "lineEnd": end_line if end_line != start else None,
                                                 "section": _section(path) or None, "paragraph": len(result.units) + 1}))
        block = []

    last = 0
    for number, line in enumerate(text.splitlines(), start=1):
        last = number
        m = RE_MD_HEADING.match(line)
        if m:
            flush(number - 1)
            level = len(m.group(1))
            headings = {k: v for k, v in headings.items() if k < level}
            headings[level] = m.group(2)
            continue
        if not line.strip():
            flush(number - 1)
        else:
            if not block:
                start = number
            block.append(line)
    flush(last)
    return result


def parse_docx(raw: bytes) -> ParseResult:
    try:
        import docx
    except ImportError as e:
        raise IngestError("DOCX support is not installed on the worker (python-docx)") from e
    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as e:
        raise IngestError(f"Could not open DOCX file: {e}")

    result = ParseResult()
    headings = {}
    for number, para in enumerate(document.paragraphs, start=1):
        body = para.text.strip()
        if not body:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        m = re.match(r"Heading (\d)", style)
        if m or style == "Title":
            level = int(m.group(1)) if m else 1
            headings = {k: v for k, v in headings.items() if k < level}
            headings[level] = body
            continue
        path = [headings[k] for k in sorted(headings)]
        result.units.append(make_unit(body, {"paragraph": number, "section": _section(path) or None}))
    for t_index, table in enumerate(document.tables, start=1):
        for r_index, row in enumerate(table.rows, start=1):
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                result.units.append(make_unit(" | ".join(cells), {"table": t_index, "row": r_index}))
    return result
